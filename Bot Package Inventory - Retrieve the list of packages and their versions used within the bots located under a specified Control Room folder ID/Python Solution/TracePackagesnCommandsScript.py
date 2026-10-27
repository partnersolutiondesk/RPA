# Note: Below external libraries to be installed for script to work.
# 1. openpyxl and requests (pip install openpyxl requests)
#
# Purpose: Walk a Control Room folder (and every subfolder, at any depth), read each taskbot and
#          report the packages and versions it uses. Output is PackageReport.xlsx in the
#          input/output folder, with two sheets:
#            "Full Report"      - one row per bot per package
#            "Bot Level Report" - one row per bot, with all of its packages summarised

import os
import json
import inspect
import logging
from collections import Counter
import requests
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment

# Global variables defined
logger = logging.getLogger(__name__)

REPORT_FILE_NAME = "PackageReport.xlsx"
FULL_REPORT_HEADER = ["Bot Id", "Bot Name", "Bot Path", "Package", "Artifact Name", "Version", "Commands Used"]
BOT_REPORT_HEADER = ["Bot Id", "Bot Name", "Bot Path", "Status", "Package Count", "Total Commands", "Packages"]


# Log the messages using logging library
def log(log_msg, log_level):
    # Automatically log the current function details.
    # Get the previous frame in the stack, otherwise it would be this function!!!
    func = inspect.currentframe().f_back.f_code

    # Dump the message + the name of this function to the log.
    logger.log(level=getattr(logging, log_level.upper(), None), msg='{0}): {1}'.format(func.co_name, log_msg))


# Initializing a logger with custom configuration
def initialize_logger(log_file_path, log_level):
    logger.setLevel(getattr(logging, log_level.upper()))
    file_handler = logging.FileHandler(log_file_path, mode='a')
    formatter = logging.Formatter('%(asctime)s %(levelname)s %(name)s (%(message)s', datefmt='(%d-%m-%Y %I.%M.%S %p)')

    file_handler.setFormatter(formatter)
    logger.addHandler(file_handler)

    log("Log file started.", 'info')


def get_report(args):
    folder_id = args[0]
    cr_url = args[1]
    username = args[2]
    apikey = args[3]
    inputoutputFolderPath = args[4]

    log_file_path = os.path.join(inputoutputFolderPath, "packagescanner.log")

    log_level = 'debug'
    initialize_logger(log_file_path=log_file_path, log_level=log_level)
    log("get_report fun called with the: folder_id: " + folder_id, log_level)
    return get_files_from_folder(folder_id, cr_url, username, apikey, inputoutputFolderPath)


# Auth function to get token
def authenticate(cr_url, username, apikey):
    auth_payload = {
        "username": username,
        "apiKey": apikey
    }
    try:
        response = requests.post(f"{cr_url}/v2/authentication", json=auth_payload)
        response.raise_for_status()
        return response.json().get("token")
    except requests.RequestException as e:
        log(f"Authentication failed: {e}", "debug")
        return None


# Calls a Control Room API. On a 401 it re-authenticates once, keeps the new token in the
# session so every later call uses it, and retries the call.
def call_api(session, method, path, **kwargs):
    url = f"{session['cr_url']}{path}"
    log(f"calling {method} {path}", "debug")
    response = requests.request(method, url, headers={"X-Authorization": str(session["token"])}, **kwargs)
    if response.status_code == 401:
        log("401 Unauthorized error. Re-authenticating for new token", "debug")
        token = authenticate(session["cr_url"], session["username"], session["apikey"])
        if not token:
            log("Re-authentication failed.", "debug")
            return None
        session["token"] = token
        log(f"calling {method} {path}", "debug")
        response = requests.request(method, url, headers={"X-Authorization": str(token)}, **kwargs)
    return response


# Returns {package name used inside bots (lower case): artifact name} for every package in the Control Room.
# Failure is logged and an empty map is returned, so the scan still runs with a blank Artifact Name column.
def get_package_artifacts(session):
    artifacts = {}
    offset = 0
    length = 100
    while True:
        payload = {
            "filterRequest": {
                "fields": [],
                "filter": None,
                "sort": [{"field": "label", "direction": "asc"}],
                "page": {"offset": offset, "length": length}
            },
            "includeDownloadUrls": False
        }
        response = call_api(session, "POST", "/v3/packages/package/list", json=payload)
        if response is None or not response.ok:
            log(f"Packages list API failed: {getattr(response, 'status_code', 'no response')}. "
                f"Artifact Name column will be blank", "debug")
            return artifacts

        items = response.json().get("list", [])
        for item in items:
            if item.get("name") and item.get("artifactName"):
                artifacts[item["name"].lower()] = item["artifactName"]

        if len(items) < length:
            break  # last page
        offset += length

    log(f"Got artifact names for {len(artifacts)} packages", "debug")
    return artifacts


def get_file_content(session, file_id):
    try:
        response = call_api(session, "GET", f"/v2/repository/files/{file_id}/content")
        if response is None:
            return None
        response.raise_for_status()
        return response.text
    except requests.RequestException as e:
        log(f"Failed to fetch content for file ID {file_id}: {e}", "debug")
        return None


# Counts every command (node) per package. Commands can be nested at any depth
# (if / loop / try blocks keep their children inside the node), so walk the whole structure.
def count_commands(obj, counts):
    if isinstance(obj, dict):
        package_name = obj.get("packageName")
        if package_name:
            counts[package_name] += 1
        for value in obj.values():
            if isinstance(value, (dict, list)):
                count_commands(value, counts)
    elif isinstance(obj, list):
        for item in obj:
            count_commands(item, counts)


# Returns a list of (package name, version, commands used) for one bot.
# Versions come from the bot's "packages" list; command counts come from the nodes.
def extract_bot_packages(content):
    versions = {}
    roots = content if isinstance(content, list) else [content]
    for root in roots:
        if not isinstance(root, dict):
            continue
        for pkg in root.get("packages", []):
            if isinstance(pkg, dict) and pkg.get("name"):
                versions[pkg["name"]] = pkg.get("version") or ""

    counts = Counter()
    count_commands(content, counts)

    # A package can be declared but unused (0 commands), or used but missing a version entry
    return [(name, versions.get(name, ""), counts.get(name, 0))
            for name in sorted(set(versions) | set(counts), key=str.lower)]


def write_report(filepath, full_rows, bot_rows):
    wb = Workbook()
    bold = Font(bold=True)

    sheet_full = wb.active
    sheet_full.title = "Full Report"
    sheet_bot = wb.create_sheet(title="Bot Level Report")

    for sheet, header, rows in ((sheet_full, FULL_REPORT_HEADER, full_rows),
                                (sheet_bot, BOT_REPORT_HEADER, bot_rows)):
        sheet.append(header)
        for cell in sheet[1]:
            cell.font = bold
        for row in rows:
            sheet.append(row)
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = sheet.dimensions

        # Size columns to content (capped), and wrap the multi-line package list
        for column in sheet.columns:
            longest = max((len(line) for cell in column for line in str(cell.value or "").split("\n")), default=10)
            sheet.column_dimensions[column[0].column_letter].width = min(longest + 2, 80)

    for row in sheet_bot.iter_rows(min_row=2, min_col=len(BOT_REPORT_HEADER), max_col=len(BOT_REPORT_HEADER)):
        for cell in row:
            cell.alignment = Alignment(wrap_text=True, vertical="top")

    wb.save(filepath)
    log(f"Report saved to {filepath}: {len(full_rows)} full report rows, {len(bot_rows)} bots", "debug")


# Main function to get files
def get_files_from_folder(folder_id, cr_url, username, apikey, inputoutputFolderPath):
    token = authenticate(cr_url, username, apikey)
    if not token:
        log("Failed to authenticate. Exiting.", "debug")
        return None

    session = {"cr_url": cr_url, "username": username, "apikey": apikey, "token": token}
    artifacts = get_package_artifacts(session)
    all_files = []
    full_rows = []
    bot_rows = []
    seen_items = set()  # Track IDs to prevent duplicates
    processed_subfolders = set()  # Track processed subfolders

    def process_bot(item):
        log("Found a taskbot " + item['id'], "debug")
        bot_info = [item['id'], item['name'], item['path']]

        content = get_file_content(session, item['id'])
        if not content:
            bot_rows.append(bot_info + ["Content not retrieved", 0, 0, ""])
            return

        try:
            packages = extract_bot_packages(json.loads(content))
        except Exception as e:
            log(f"JSON parse error for {item.get('name')}: {e}", "debug")
            bot_rows.append(bot_info + ["Parse error", 0, 0, ""])
            return

        for name, version, commands in packages:
            full_rows.append(bot_info + [name, artifacts.get(name.lower(), ""), version, commands])
        package_lines = "\n".join(f"{name} ({version or 'version unknown'})" for name, version, _ in packages)
        bot_rows.append(bot_info + ["OK", len(packages), sum(commands for _, _, commands in packages), package_lines])

    def fetch_folder(folder_id):
        log("inside fetch folder " + folder_id, "debug")
        page_number = 0

        while True:
            payload = {
                "filter": {},
                "sort": [{"field": "id", "direction": "desc"}],
                "page": {
                    "offset": page_number * 100,
                    "length": 100
                }
            }
            response = call_api(session, "POST", f"/v2/repository/folders/{folder_id}/list", json=payload)
            if response is None or not response.ok:
                log(f"Listing folder {folder_id} failed: {getattr(response, 'status_code', 'no response')}", "debug")
                return

            items = response.json().get("list", [])

            # Stop pagination if no new items are received
            new_items = [item for item in items if item['id'] not in seen_items]
            if not new_items:
                break

            for item in new_items:
                seen_items.add(item['id'])  # Track item ID to avoid duplication

                if item['type'] == 'application/vnd.aa.taskbot':
                    all_files.append(item)  # Store full item details
                    process_bot(item)

                elif item['type'] == 'application/vnd.aa.directory' and item['id'] not in processed_subfolders:
                    processed_subfolders.add(item['id'])  # Mark subfolder as processed
                    fetch_folder(item['id'])  # Recursively fetch subfolder contents

            page_number += 1  # Move to the next page for pagination

    try:
        fetch_folder(folder_id)
    finally:
        # Save whatever was collected, even if the run was interrupted part way
        full_rows.sort(key=lambda r: (str(r[2]).lower(), str(r[1]).lower(), str(r[3]).lower()))
        bot_rows.sort(key=lambda r: (str(r[2]).lower(), str(r[1]).lower()))
        write_report(os.path.join(inputoutputFolderPath, REPORT_FILE_NAME), full_rows, bot_rows)

    return json.dumps({"list": all_files}, indent=4)


# Configure the five values below before running the script.
# Do not commit real credentials to source control.
args = [
    'REPLACE_WITH_FOLDER_ID',                     # 1. ID of the folder to scan
    'https://your-control-room.example.com',      # 2. Control Room URL, without a trailing slash
    'REPLACE_WITH_USERNAME',                      # 3. User name
    'REPLACE_WITH_API_KEY',                       # 4. API key
    r'C:\Reports\PackageScan',                    # 5. Output folder (must exist and be writable)
]
get_report(args)
